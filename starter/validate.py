"""Проверка ГОТОВОГО документа перед выдачей. Детерминированная, без AI.

Ключевая идея: проверяем не форму, а РЕЗУЛЬТАТ — тот самый .docx, который
скачает оператор. Форма может быть заполнена правильно, а документ всё равно
выйдет битым: не отработала подстановка, поехала таблица, шаблон не тот.

Разделение строгое и оно важно:
  errors   — документ отдавать НЕЛЬЗЯ (незакрытые плейсхолдеры, пустые
             обязательные поля, сумма ≤ 0);
  warnings — подозрительно, но решает человек (контрольные суммы, пустой
             состав услуг).

Чего эта проверка НЕ ловит: шрифт и вёрстку. Ни один текстовый тест не увидит,
что документ собрался в Calibri вместо Times New Roman или что таблица уехала
на вторую страницу. Это ловится только глазами — см. docs/06-proverka-i-testy.md.
"""
import re

from docx import Document

from requisites import check_requisites

# Поля, без которых документ не документ. Ключ → как назвать оператору.
REQUIRED_FIELDS = {
    "customer_name": "Наименование контрагента",
    "customer_inn_kpp": "ИНН контрагента",
    "customer_legal_addr": "Юридический адрес контрагента",
    "num": "№ документа",
    "doc_date": "Дата документа",
    "total": "Сумма",
}

PLACEHOLDER_RE = re.compile(r"\{\{|\}\}|\{%")


def iter_text(doc):
    """Весь текст документа: параграфы + ячейки таблиц + колонтитулы.

    Обход именно такой, потому что плейсхолдеры любят прятаться там, куда
    не смотрят: в шапке таблицы реквизитов и в колонтитуле с номером договора.
    """
    for p in doc.paragraphs:
        yield p.text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p.text
    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                yield p.text


def check_document(docx_path, form):
    """Готовый .docx + исходная форма → {'errors': [...], 'warnings': [...]}."""
    errors, warnings = [], []

    for field, label in REQUIRED_FIELDS.items():
        if not str(form.get(field) or "").strip():
            errors.append("Не заполнено: {}.".format(label))

    try:
        if int(str(form.get("total") or 0)) <= 0:
            errors.append("Сумма должна быть больше нуля.")
    except (ValueError, TypeError):
        errors.append("Сумма указана некорректно.")

    doc = Document(str(docx_path))
    texts = list(iter_text(doc))

    leftovers = [t.strip() for t in texts if PLACEHOLDER_RE.search(t or "")]
    if leftovers:
        errors.append(
            "В документе остались неподставленные плейсхолдеры ({} шт.): {}".format(
                len(leftovers), leftovers[0][:120]
            )
        )

    # Пустые прочерки на месте реквизитов — верный признак, что имя поля
    # в шаблоне и в контексте разошлось. Плейсхолдер при этом НЕ остаётся:
    # docxtpl молча подставляет пустую строку. Ловим по факту пустоты.
    body = "\n".join(texts)
    for field, label in (("customer_name", "наименование контрагента"),
                         ("customer_inn_kpp", "ИНН контрагента")):
        value = str(form.get(field) or "").strip()
        if value and value not in body:
            errors.append(
                "В документе нет значения поля «{}» — вероятно, имя поля в шаблоне "
                "не совпадает с ключом контекста.".format(label)
            )

    warnings.extend(check_requisites(form))
    if not any("- " in (t or "") for t in texts):
        warnings.append("В документе не найдено ни одного пункта состава услуг.")

    return {"errors": errors, "warnings": warnings}
