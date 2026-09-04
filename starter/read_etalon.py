"""Прочитать .docx текстом — чтобы отдать эталон AI-агенту.

Зачем отдельная утилита. Файл .docx — это zip с XML внутри, а не текст.
AI-агент не может просто «открыть» его: он увидит бинарный мусор или откажется
читать вовсе. Поэтому эталон сначала извлекается в текст, и уже текст уходит
в чат или в контекст агента.

Абзацы печатаются С НОМЕРАМИ. Это не украшение: build-скрипт шаблона находит
абзацы ПО ТЕКСТУ и переписывает их целиком (см. docs/02-shablon-docx.md), и
номера — общий язык, на котором вы с агентом договариваетесь, какой именно
абзац становится плейсхолдером.

Использование:

    # разобрать свой договор перед сборкой шаблона
    python read_etalon.py references/moy-dogovor.docx

    # проверить, какие плейсхолдеры реально попали в готовый шаблон
    python read_etalon.py templates/dogovor_services.docx --placeholders

    # сохранить в файл, чтобы приложить к чату
    python read_etalon.py references/moy-dogovor.docx > /tmp/etalon.txt
"""
import re
import sys
from pathlib import Path

from docx import Document

# Поля вида {{ name }}
FIELD_RE = re.compile(r"\{\{\s*([\w.]+)")
# Управляющие теги вида {%p for item in services %}. Из них полем является
# КОЛЛЕКЦИЯ (services), а не переменная цикла (item) и не само слово for:
# в словаре полей есть «состав услуг», а «item» и «endfor» там взяться неоткуда
# и только сбивают при сверке.
LOOP_RE = re.compile(r"\{%-?\s*p?\s*for\s+\w+\s+in\s+([\w.]+)")


def paragraphs_with_index(doc):
    """Непустые абзацы тела документа с их порядковыми номерами."""
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            yield i, text


def dump_text(path):
    doc = Document(str(path))

    print("# ЭТАЛОН: {}".format(Path(path).name))
    print("#")
    print("# Абзацы пронумерованы — ссылайся на номера, когда договариваешься")
    print("# с агентом, какое место документа становится полем.")
    print()

    print("## АБЗАЦЫ")
    for i, text in paragraphs_with_index(doc):
        print("[{:>3}] {}".format(i, text))

    if doc.tables:
        print()
        print("## ТАБЛИЦЫ")
        # Реквизиты сторон почти всегда живут в таблице, и именно там прячется
        # половина переменных полей. Без явного обхода таблиц агент их не увидит.
        for t_i, table in enumerate(doc.tables):
            print()
            print("[таблица {}] {}×{}".format(t_i, len(table.rows), len(table.columns)))
            for r_i, row in enumerate(table.rows):
                cells = [c.text.strip().replace("\n", " ¶ ") for c in row.cells]
                print("  ({},{}) {}".format(t_i, r_i, " | ".join(cells)))

    sections_with_text = [
        (kind, p.text.strip())
        for s in doc.sections
        for kind, part in (("колонтитул-верх", s.header), ("колонтитул-низ", s.footer))
        for p in part.paragraphs
        if p.text.strip()
    ]
    if sections_with_text:
        print()
        print("## КОЛОНТИТУЛЫ")
        for kind, text in sections_with_text:
            print("  [{}] {}".format(kind, text))


def dump_placeholders(path):
    """Все плейсхолдеры готового шаблона — сверить со словарём полей.

    Две вещи, которые эта проверка ловит сразу: опечатку в имени поля
    (`{{ customer_nmae }}` попадёт в список и не совпадёт со словарём) и
    плейсхолдер, спрятавшийся в таблице или колонтитуле.
    """
    doc = Document(str(path))
    found = {}
    loops = {}

    def scan(text, where):
        for m in FIELD_RE.finditer(text or ""):
            found.setdefault(m.group(1), []).append(where)
        for m in LOOP_RE.finditer(text or ""):
            loops.setdefault(m.group(1), []).append(where)

    for i, text in paragraphs_with_index(doc):
        scan(text, "абзац {}".format(i))
    for t_i, table in enumerate(doc.tables):
        for r_i, row in enumerate(table.rows):
            for c_i, cell in enumerate(row.cells):
                scan(cell.text, "таблица {} ({},{})".format(t_i, r_i, c_i))
    for s in doc.sections:
        for kind, part in (("верх", s.header), ("низ", s.footer)):
            for p in part.paragraphs:
                scan(p.text, "колонтитул-" + kind)

    # Переменная цикла ({{ item }}) — не поле словаря, а имя внутри цикла.
    loop_vars = set()
    for text_source in (doc.paragraphs,):
        for p in text_source:
            m = re.search(r"\{%-?\s*p?\s*for\s+(\w+)\s+in\s+", p.text or "")
            if m:
                loop_vars.add(m.group(1))
    for var in loop_vars:
        found.pop(var, None)

    print("# ПЛЕЙСХОЛДЕРЫ: {}".format(Path(path).name))
    print("# Всего полей: {}".format(len(found)))
    print()
    for name in sorted(found):
        places = found[name]
        print("{:<28} {}× — {}".format(name, len(places), ", ".join(places[:4])))
    if loops:
        print()
        print("## СПИСКИ (цикл по коллекции)")
        for name in sorted(loops):
            print("{:<28} {}".format(name, ", ".join(loops[name][:4])))
    if not found:
        print("Плейсхолдеров нет. Это эталон, а не шаблон — либо подстановка")
        print("не найдёт ни одного поля (проверь, не набраны ли скобки вручную:")
        print("см. docs/02-shablon-docx.md, Word дробит текст на runs).")


def main():
    args = [a for a in sys.argv[1:]]
    if not args:
        print(__doc__)
        return 1
    path = Path(args[0])
    if not path.exists():
        print("Нет файла: {}".format(path))
        return 1
    if "--placeholders" in args:
        dump_placeholders(path)
    else:
        dump_text(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
